from __future__ import annotations

import fitz
import io
from pathlib import Path
from openpyxl import Workbook

from models.plan_profile import PlanProfile
from services.file_order import list_file_meta, order_uploads
from services.orchestrator import _split_summaries
from services.pdf_pages import pdf_page_count, pdf_page_texts
from services.plan_profile_service import finalize_profile
from services.assembly_log import _allocation
from services.rules_engine import _conditions_met
from services.plan_config import rules_for
from services.source_profile import parse_excel
from services.source_profile import parse_summary_text

ROOT = Path(__file__).resolve().parents[1]


def make_text_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((40, 60), text)
    data = doc.tobytes()
    doc.close()
    return data


def test_autorder_prioritizes_result_summary_before_action_and_reports():
    uploads = [
        ("901234 2025 Top Heavy.pdf", b"th"),
        ("901234 2025 Action Required.pdf", b"action"),
        ("901234 2025 ResultSumm.pdf", b"summary"),
        ("901234 2025 401k Valuation Pkg.pdf", b"cover"),
        ("901234 2025 Contri.pdf", b"contribution"),
        ("901234 2025 HCEKey.pdf", b"hce"),
    ]

    result = [name for name, _ in order_uploads(uploads)]

    assert result == [
        "901234 2025 401k Valuation Pkg.pdf",
        "901234 2025 ResultSumm.pdf",
        "901234 2025 Action Required.pdf",
        "901234 2025 Contri.pdf",
        "901234 2025 HCEKey.pdf",
        "901234 2025 Top Heavy.pdf",
    ]


def test_generic_result_summary_is_included_in_assembled_pdf():
    generic_summary = make_text_pdf(
        "Compliance Testing Summary of Results\nPlan Number: 901234\nADP Test: Pass"
    )

    valuation, summaries = _split_summaries([("Results.pdf", generic_summary)])

    assert [name for name, _ in valuation] == ["Results.pdf"]
    assert [name for name, _ in summaries] == ["Results.pdf"]


def test_valuation_package_with_summary_front_page_is_not_duplicated():
    package = make_text_pdf(
        "Compliance Testing Summary of Results\nPlan Number: 901234\nADP Test: Pass"
    )

    valuation, summaries = _split_summaries([("901234 2025 401k Valuation Pkg.pdf", package)])

    assert [name for name, _ in valuation] == ["901234 2025 401k Valuation Pkg.pdf"]
    assert summaries == []


def test_403b_packet_pages_are_labeled_from_their_actual_headings():
    path = ROOT / "sample-data/403b/pass/input-pdfs/555403 2025 403b Valuation Pkg.pdf"
    data = path.read_bytes()
    meta = list_file_meta(path.name, len(data), pdf_page_count(data), pdf_page_texts(data))

    assert meta["labels"] == [
        "Cover Letter",
        "Action required",
        "Action required",
        "Action required",
        "ADP/ACP Failure excess page after 12 months (Current year Testing Method)",
        "ADP/ACP Failure excess page after 12 months (Prior year Testing Method)",
        "ADP/ACP Failure excess page Current year",
        "415 Failure information",
        "ADP/ACP Failure letter",
        "402g Failure letter",
        "415 Failure letter",
        "Year End Recap",
    ]


def test_summary_extraction_does_not_read_failure_examples_after_summary_section():
    text = (
        "Compliance Testing Summary of Results\n"
        "Plan Year End: December 31, 2025\n"
        "GRACE MEDICAL HOME, INC. 403(B) PLAN\n"
        "555403\n"
        "Average Contribution Percentage (ACP)\nTested: Pass\n"
        "Annual Valuation Report\n"
        "IMMEDIATE ACTION REQUIRED - ACP TEST FAILURE\n"
        "Contribution refunds to participants are required.\n"
        "Failed Compliance Testing - After 12-months\n"
    )

    parsed = parse_summary_text(text)

    assert parsed["acp_failed"] is False
    assert "returns_required" not in parsed
    assert parsed["after_12_months"] is False


def test_unknown_negative_rule_fact_is_skipped_not_assumed_false():
    profile = PlanProfile(hce_current=None, hce_future=None)

    matched, skipped, reason = _conditions_met(
        {"all": [{"field": "hce_current", "not_true": True}, {"field": "hce_future", "not_true": True}]},
        profile,
    )

    assert not matched and skipped
    assert "missing hce_current" in reason
    assert "missing hce_future" in reason


def test_top_heavy_threshold_is_strictly_greater_than_sixty():
    at_threshold = PlanProfile(top_heavy_percent=60.0)
    above_threshold = PlanProfile(top_heavy_percent=60.01)

    assert finalize_profile(at_threshold).top_heavy is False
    assert finalize_profile(above_threshold).top_heavy is True


def _excel_with_rows(rows: list[tuple[str, str]]) -> bytes:
    book = Workbook()
    sheet = book.active
    for index, (label, value) in enumerate(rows, start=5):
        sheet.cell(index, 1, label)
        sheet.cell(index, 2, value)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_annual_allocation_is_not_misclassified_as_explicit_true_up():
    parsed = parse_excel(_excel_with_rows([
        ("Annual allocation (per document including true up)? (yes / no)", "Yes"),
    ]))

    assert parsed["contributions_required"] is True
    assert parsed["true_up"] is False
    profile = PlanProfile(**parsed)
    assert _allocation(profile) == ("Annual", "")


def test_explicit_true_up_requires_contributions_and_variance_copy():
    parsed = parse_excel(_excel_with_rows([("True-up", "Yes")]))

    assert parsed["contributions_required"] is True
    assert parsed["true_up"] is True
    profile = PlanProfile(**parsed)
    assert _allocation(profile) == ("Variance", "0.00 variance")


def test_403b_after_twelve_month_wording_requires_outstanding_returns():
    current_rule = next(rule for rule in rules_for("403b")["rules"] if rule["id"] == "D")
    profile = PlanProfile(testing_method="CURRENT", after_12_months=True, returns_required=False, acp_failed=True)

    matched, skipped, _reason = _conditions_met(current_rule["keep_when"], profile)

    assert not matched
    assert not skipped
